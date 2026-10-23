# -*- coding: utf-8 -*-
"""Người xây dựng: Hoàng Đức Vũ – Trưởng phòng Kỹ thuật, Công ty CP TEXO Tư vấn và Đầu tư.
Tạo đề cương TVGS theo khung chuẩn TEXO – mẫu Cơ bản hoặc Nâng cao, phù hợp quy mô công trình."""
import copy
import datetime as dt
import json
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

from . import texo_docx_kit as K

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
FONT = "Times New Roman"
BODY = 13


def load(name):
    return json.load(open(os.path.join(DATA, name), encoding="utf-8"))


# =====================================================================  GỢI Ý THEO QUY MÔ
LOAI_CT = ["Dân dụng – nhà cao tầng", "Dân dụng – nhà thấp tầng, công trình công cộng", "Công nghiệp – nhà xưởng, kho",
           "Hạ tầng kỹ thuật, giao thông nội bộ", "Công trình xử lý nước thải, hạ tầng môi trường", "Gói chuyên đề – thi công cọc", "Gói chuyên đề – khoan khảo sát địa chất",
           "Gói chuyên đề – trắc đạc, quan trắc"]
CAP_CT = ["Đặc biệt", "I", "II", "III", "IV"]
MONG = ["Cọc ép / cọc đóng BTCT", "Cọc khoan nhồi", "Tường vây / Barrette", "Móng nông"]
KET_CAU = ["BTCT toàn khối", "Cáp dự ứng lực", "Kết cấu thép"]


def de_xuat_muc_mau(info):
    """Trả (mức mẫu, lý do)."""
    loai = info.get("loai_ct", "")
    if loai.startswith("Gói chuyên đề"):
        return "Cơ bản", "Gói thầu chuyên đề một công tác – dùng mẫu Cơ bản."
    ly_do = []
    cap = info.get("cap_ct", "III")
    if cap in ("Đặc biệt", "I"):
        ly_do.append(f"công trình cấp {cap}")
    if int(info.get("so_ham") or 0) >= 1:
        ly_do.append(f"{info.get('so_ham')} tầng hầm")
    if int(info.get("so_tang") or 0) >= 9:
        ly_do.append(f"{info.get('so_tang')} tầng nổi")
    if float(info.get("dt_san") or 0) >= 20000:
        ly_do.append("tổng diện tích sàn ≥ 20.000 m²")
    if info.get("cdt_yeu_cau_cao"):
        ly_do.append("Chủ đầu tư yêu cầu hệ thống quản lý chặt (quy trình, ma trận, ITP)")
    if ly_do:
        return "Nâng cao", "Đề xuất mẫu Nâng cao vì: " + "; ".join(ly_do) + "."
    return "Cơ bản", "Quy mô vừa và nhỏ – mẫu Cơ bản là đủ; có thể chọn Nâng cao nếu Chủ đầu tư yêu cầu."


def de_xuat_hang_muc(info):
    loai = info.get("loai_ct", "")
    so_tang = int(info.get("so_tang") or 0)
    so_ham = int(info.get("so_ham") or 0)
    cap = info.get("cap_ct", "III")
    mong = info.get("mong", [])
    kc = info.get("ket_cau", [])
    s = ["NANGLUC", "VATLIEU"]
    if loai == "Gói chuyên đề – thi công cọc":
        return s + ["TRACDIA", "EPCOC" if "Cọc khoan nhồi" not in mong else "KHOANNHOI", "TNCOC", "ATLD"]
    if loai == "Gói chuyên đề – khoan khảo sát địa chất":
        return ["NANGLUC", "KHOANKS", "ATLD"]
    if loai == "Gói chuyên đề – trắc đạc, quan trắc":
        return ["NANGLUC", "TRACDAC", "ATLD"]
    s += ["TRACDIA"]
    if "Cọc ép / cọc đóng BTCT" in mong:
        s += ["EPCOC", "TNCOC"]
    if "Cọc khoan nhồi" in mong:
        s += ["KHOANNHOI", "TNCOC"]
    if "Tường vây / Barrette" in mong:
        s += ["TUONGVAY"]
    s += ["DAODAT", "MONG"]
    if so_ham:
        s += ["HAM"]
    if so_ham or so_tang >= 15 or "Tường vây / Barrette" in mong:
        s += ["QUANTRAC"]
    if loai.startswith("Công trình xử lý nước thải"):
        s = ["NANGLUC", "VATLIEU", "TRACDIA", "DAODAT"] + [x for x in s if x in ("EPCOC", "KHOANNHOI", "TNCOC")] + \
            ["MONG", "BE_BTCT", "BE_THAM", "XAYTRAT", "SONBA", "CHONGTHAM", "ONG_CN", "TB_CN", "DIEN_DK", "CHONGSET", "CTN", "PCCC",
             "HTKT", "CAYXANH", "AT_HC", "CHAYTHU_CN", "ITP_CN", "ATLD"]
        return list(dict.fromkeys(s))
    if loai.startswith("Hạ tầng"):
        s = ["NANGLUC", "VATLIEU", "TRACDIA", "DAODAT", "HTKT", "CTN", "CSNGOAI", "RANHCAP", "CAYXANH", "ATLD"]
        return list(dict.fromkeys(s))
    if "BTCT toàn khối" in kc or not kc:
        s += ["BTCT"]
    if "Cáp dự ứng lực" in kc:
        s += ["DUL"]
    if "Kết cấu thép" in kc or loai.startswith("Công nghiệp"):
        s += ["KCTHEP"]
    s += ["CHONGTHAM", "XAYTRAT", "CUANHOM", "SONBA", "TRAN", "CHONGMOI"]
    if loai.startswith("Dân dụng"):
        s += ["CUAGO", "CUACC", "TBVS"]
        if info.get("noi_that"):
            s += ["SANGO", "NOITHAT", "LANCAN"]
    if loai.startswith("Công nghiệp"):
        s += ["EPOXY", "CUACC"]
    s += ["MEPCHUNG", "ONGDIEN", "MANGCAP", "CAPDIEN", "TBDIEN", "CHONGSET", "CTN", "PCCC", "DIENNHE"]
    if loai.startswith("Công nghiệp") or cap in ("Đặc biệt", "I", "II") or so_tang >= 5:
        s += ["HVAC"]
    if so_tang >= 6 or info.get("thang_may"):
        s += ["THANGMAY"]
    if cap in ("Đặc biệt", "I", "II") or loai.startswith("Công nghiệp") or info.get("tba"):
        s += ["TBA", "MAYPHAT"]
    if so_tang >= 15 or cap in ("Đặc biệt", "I"):
        s += ["BUSWAY", "BMS"]
    if so_ham:
        s += ["KSXE", "EPOXY"]
    if info.get("xlnt"):
        s += ["XLNT"]
    if info.get("ha_tang_ngoai_nha", True):
        s += ["HTKT", "CSNGOAI", "RANHCAP", "CAYXANH"]
    s += ["CHAYTHU", "ATLD"]
    return list(dict.fromkeys(s))


def de_xuat_nhan_su(info, modules, level):
    cap = info.get("cap_ct", "III")
    n_xd = {"Đặc biệt": 4, "I": 3, "II": 2, "III": 1, "IV": 1}.get(cap, 1)
    if float(info.get("dt_san") or 0) > 40000:
        n_xd += 1
    rows = [dict(chuc_danh="Tư vấn giám sát trưởng", ho_ten=info.get("tvgs_truong", ""), chuyen_mon="Kỹ sư xây dựng; chứng chỉ hành nghề giám sát phù hợp hạng công trình",
                 nhiem_vu="Điều hành đoàn TVGS; đầu mối làm việc với CĐT; ký biên bản nghiệm thu giai đoạn, hoàn công, xác nhận khối lượng, báo cáo")]
    loai = info.get("loai_ct", "")
    if loai.startswith("Gói chuyên đề"):
        rows.append(dict(chuc_danh="Kỹ sư giám sát chuyên ngành", ho_ten="", chuyen_mon="Phù hợp công tác của gói thầu", nhiem_vu="Giám sát, nghiệm thu công việc; ghi nhật ký; lập báo cáo"))
    else:
        for k in range(n_xd):
            rows.append(dict(chuc_danh=f"Kỹ sư giám sát xây dựng {k + 1}" if n_xd > 1 else "Kỹ sư giám sát xây dựng", ho_ten="",
                             chuyen_mon="Kỹ sư xây dựng DD&CN; chứng chỉ giám sát lĩnh vực dân dụng – công nghiệp",
                             nhiem_vu="Giám sát nền móng, kết cấu, hoàn thiện theo phân công; nghiệm thu công việc; nhật ký"))
        mep = [m for m in modules if m in ("MEPCHUNG", "ONGDIEN", "CAPDIEN", "TBDIEN", "CTN", "PCCC", "HVAC", "THANGMAY", "TBA")]
        if mep:
            n_mep = 2 if cap in ("Đặc biệt", "I") else 1
            for k in range(n_mep):
                rows.append(dict(chuc_danh=f"Kỹ sư giám sát cơ điện {k + 1}" if n_mep > 1 else "Kỹ sư giám sát cơ điện (MEP)", ho_ten="",
                                 chuyen_mon="Kỹ sư điện / cấp thoát nước / nhiệt; chứng chỉ giám sát lắp đặt thiết bị",
                                 nhiem_vu="Giám sát lắp đặt điện, nước, PCCC, ĐHKK, thang máy; chứng kiến thử nghiệm, chạy thử"))
        if any(m in modules for m in ("HTKT", "CAYXANH")) and cap in ("Đặc biệt", "I", "II"):
            rows.append(dict(chuc_danh="Kỹ sư giám sát hạ tầng", ho_ten="", chuyen_mon="Kỹ sư cầu đường / cấp thoát nước",
                             nhiem_vu="Giám sát hạ tầng kỹ thuật ngoài nhà, đường nội bộ, cảnh quan"))
        if level == "Nâng cao":
            rows.append(dict(chuc_danh="Kỹ sư QA-QC / quản lý hồ sơ, khối lượng", ho_ten="", chuyen_mon="Kỹ sư xây dựng / kinh tế xây dựng",
                             nhiem_vu="Quản lý hồ sơ, danh mục trình duyệt, NCR/RFI; kiểm tra hồ sơ khối lượng"))
            rows.append(dict(chuc_danh="Giám sát trắc đạc", ho_ten="", chuyen_mon="Kỹ sư trắc địa", nhiem_vu="Kiểm tra mốc, tọa độ, cao độ, độ thẳng đứng, hoàn công, quan trắc"))
    rows.append(dict(chuc_danh="Giám sát an toàn lao động" + (" (kiêm nhiệm)" if level == "Cơ bản" else ""), ho_ten="",
                     chuyen_mon="Chứng nhận huấn luyện ATVSLĐ phù hợp", nhiem_vu="Giám sát ATLĐ, VSMT, PCCC tại công trường; kiểm tra hồ sơ an toàn của nhà thầu"))
    return rows


# =====================================================================  TIỆN ÍCH DỰNG WORD
class Doc:
    def __init__(self, info, level):
        self.info, self.level = info, level
        self.d = Document()
        self.n = [0, 0, 0, 0]          # PHẦN, 1., 1.1., 1.1.1.
        self.heads = []
        self._styles()

    # ---- kiểu chữ
    def _styles(self):
        st = self.d.styles
        normal = st["Normal"]
        normal.font.name, normal.font.size = FONT, Pt(BODY)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        spec = {1: (14, "0C0150", 12, 6), 2: (13.5, "1F3864", 10, 4), 3: (13, "2E5496", 8, 3), 4: (13, "3F3F3F", 6, 2)}
        for lv, (sz, col, b4, af) in spec.items():
            s = st[f"Heading {lv}"]
            s.font.name, s.font.size, s.font.bold = FONT, Pt(sz), True
            s.font.italic = False
            s.font.color.rgb = RGBColor.from_string(col)
            rpr = s.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.insert(0, rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            pf = s.paragraph_format
            pf.space_before, pf.space_after, pf.keep_with_next = Pt(b4), Pt(af), True
            pf.line_spacing = 1.15

    # ---- tiêu đề có đánh số
    def part(self, title):
        self.n[0] += 1
        roman = ["I", "II", "III", "IV", "V", "VI"][self.n[0] - 1]
        txt = f"PHẦN {roman}. {title.upper()}"
        p = self.d.add_paragraph(txt, style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.heads.append((1, txt))
        return p

    def h(self, lv, title):
        """lv 2 → '1.', lv 3 → '1.1.', lv 4 → '1.1.1.'"""
        i = lv - 1
        self.n[i] += 1
        for j in range(i + 1, 4):
            self.n[j] = 0
        num = ".".join(str(x) for x in self.n[1:i + 1]) + "."
        txt = f"{num} {title}"
        self.d.add_paragraph(txt, style=f"Heading {lv}")
        self.heads.append((lv, txt))

    def sub(self, title):
        """Tiêu đề nhỏ không đánh số (in đậm nghiêng) – không vào mục lục."""
        p = self.p("", after=2, keep=True)
        K.run(p, title, BODY, b=True, i=True, c="1F3864")
        return p

    def p(self, text="", after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY, b=False, i=False, keep=False, size=BODY, indent=None):
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing, pf.keep_with_next = Pt(0), Pt(after), 1.2, keep
        if indent is not None:
            pf.first_line_indent = Cm(indent)
        par.alignment = align
        if text:
            K.run(par, self.fill(text), size, b=b, i=i)
        return par

    def b(self, text, level=0, prefix="-"):
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(3), 1.2
        left = 0.63 + 0.6 * level
        pf.left_indent, pf.first_line_indent = Cm(left), Cm(-0.45)
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        K.run(par, (prefix if level == 0 else "+") + " ", BODY)
        K.run(par, self.fill(text), BODY)
        return par

    def bullets(self, items, level=0):
        for x in items:
            self.b(x, level)

    def fill(self, text):
        f = self.info
        rep = {"{DU_AN}": f.get("du_an") or "dự án", "{CDT}": f.get("chu_dau_tu") or "Chủ đầu tư",
               "{THOI_GIAN}": str(f.get("thoi_gian") or "…"), "{CONG_TRINH}": f.get("cong_trinh") or "công trình"}
        for k, v in rep.items():
            text = text.replace(k, v)
        return text

    def table(self, widths, head, rows, size=11, center_cols=(), sev_col=None):
        t = K.data_table(self.d, widths, head, [[self.fill(c) if isinstance(c, str) else c for c in r] for r in rows],
                         size=size, center_cols=center_cols, sev_col=sev_col)
        self.p("", after=4)
        return t

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- khối đoạn lấy từ thư viện nội dung (mẫu Nâng cao)
    def snippet(self, items, width=9026, skip_heads=(), h4=True):
        for kind, val in items:
            if kind == "t":
                rows = val
                if not rows:
                    continue
                ncol = max(len(r) for r in rows)
                rows = [r + [""] * (ncol - len(r)) for r in rows]
                head, body = rows[0], rows[1:]
                if not body:
                    continue
                first = 0
                if ncol == 2:
                    widths = [2600, width - 2600]
                else:
                    first = 700 if re.fullmatch(r"(TT|STT|Bước)", head[0].strip()) else 2000
                    rest = (width - first) // (ncol - 1)
                    widths = [first] + [rest] * (ncol - 1)
                    widths[-1] += width - sum(widths)
                self.table(widths, head, body, size=10.5, center_cols=(0,) if first == 700 else ())
            elif kind == "h":
                t = re.sub(r"^(\d+(\.\d+)*\.?|[a-g]\.)\s*", "", val).strip()
                if t and t not in skip_heads:
                    if len(t) > 110:
                        self.p(t)
                    elif h4 and self.n[2] > 0:
                        self.h(4, t.rstrip(":"))
                    else:
                        self.sub(t.rstrip(":") + ":" if not t.endswith(":") else t)
            elif kind == "b":
                self.b(val)
            else:
                self.p(val)


# =====================================================================  CÁC PHẦN CỦA ĐỀ CƯƠNG
def _cover(D, approval=False):
    f = D.info
    d = D.d
    if os.path.exists(K.LOGO):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(K.LOGO, width=Cm(5.2))
    D.p("CÔNG TY CỔ PHẦN TEXO TƯ VẤN VÀ ĐẦU TƯ", align=WD_ALIGN_PARAGRAPH.CENTER, b=True, size=13, after=2)
    D.p("", after=30)  # tài liệu phát hành ra ngoài chỉ mang tên Công ty, không nêu đơn vị nội bộ
    p = D.p("", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    K.run(p, "ĐỀ CƯƠNG", 26, b=True, c="0C0150")
    p = D.p("", align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    K.run(p, "TƯ VẤN GIÁM SÁT THI CÔNG XÂY DỰNG", 18, b=True, c="0C0150")
    pairs = [("DỰ ÁN", f.get("du_an")), ("CÔNG TRÌNH", f.get("cong_trinh")), ("GÓI THẦU", f.get("goi_thau")),
             ("ĐỊA ĐIỂM", f.get("dia_diem")), ("CHỦ ĐẦU TƯ", f.get("chu_dau_tu")), ("HỢP ĐỒNG", f.get("hop_dong")),
             ("TƯ VẤN GIÁM SÁT", "CÔNG TY CỔ PHẦN TEXO TƯ VẤN VÀ ĐẦU TƯ")]
    pairs = [(a, b) for a, b in pairs if b]
    t = K.new_table(d, [2600, 400, 6026], rows=len(pairs))
    for (a, b), r in zip(pairs, t.rows):
        for c in r.cells:
            K._tc_borders(c, {})
            K.margins(c, 50, 50)
        K.cp(r.cells[0], a, 12.5, b=True)
        K.cp(r.cells[1], ":", 12.5, b=True)
        K.cp(r.cells[2], b.upper() if a in ("DỰ ÁN", "CÔNG TRÌNH", "GÓI THẦU") else b, 12.5, b=a in ("DỰ ÁN", "CÔNG TRÌNH", "GÓI THẦU"))
    if approval:
        D.p("", after=36)
        t = K.new_table(d, [4513, 4513], rows=4)
        for r in t.rows:
            for c in r.cells:
                K._tc_borders(c, {})
        K.cp(t.cell(0, 0), "CHỦ ĐẦU TƯ", 12.5, b=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        K.cp(t.cell(1, 0), "PHÊ DUYỆT", 12.5, b=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        K.cp(t.cell(0, 1), "ĐẠI DIỆN TƯ VẤN GIÁM SÁT", 12.5, b=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        K.cp(t.cell(1, 1), "CÔNG TY CỔ PHẦN TEXO TƯ VẤN VÀ ĐẦU TƯ", 11.5, b=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        for k in (0, 1):
            K.cp(t.cell(2, k), "(Ký, ghi rõ họ tên, đóng dấu)", 11, i=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=70)
        K.cp(t.cell(3, 1), f.get("nguoi_ky_tvgs", ""), 12.5, b=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        D.p("", after=90)
    ngay = f.get("ngay") or dt.date.today()
    D.p(f"{(f.get('noi_lap') or 'Hà Nội').upper()}, THÁNG {ngay.month:02d}/{ngay.year}", align=WD_ALIGN_PARAGRAPH.CENTER, b=True, size=12.5)
    D.page_break()


def _toc(D):
    p = D.p("MỤC LỤC", align=WD_ALIGN_PARAGRAPH.CENTER, b=True, size=14, after=8)
    par = D.d.add_paragraph()
    r = par.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = 'TOC \\o "1-3" \\h \\z \\u'
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    r._r.append(a); r._r.append(i); r._r.append(s)
    D._toc_par = par
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    D._toc_end = e
    return par


def _fill_toc(D):
    """Ghi sẵn danh sách tiêu đề vào vùng kết quả của trường TOC (Word cập nhật số trang khi mở)."""
    par = D._toc_par
    anchor = par
    for lv, txt in D.heads:
        if lv > 3:
            continue
        np_ = copy.deepcopy(par._p)
        for ch in list(np_):
            if ch.tag != qn("w:pPr"):
                np_.remove(ch)
        anchor._p.addnext(np_)
        from docx.text.paragraph import Paragraph
        newp = Paragraph(np_, par._parent)
        newp.paragraph_format.left_indent = Cm(0.6 * (lv - 1))
        newp.paragraph_format.space_after = Pt(1)
        K.run(newp, txt, 12 if lv > 1 else 12.5, b=(lv == 1))
        anchor = newp
    r = anchor.add_run()
    r._r.append(D._toc_end)
    # yêu cầu Word cập nhật trường khi mở
    settings = D.d.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def _viet_tat(D):
    D.p("TỪ NGỮ VIẾT TẮT", align=WD_ALIGN_PARAGRAPH.CENTER, b=True, size=13.5, after=6)
    rows = [["CĐT", "Chủ đầu tư"], ["BQLDA", "Ban quản lý dự án / đại diện Chủ đầu tư"], ["TVGS", "Tư vấn giám sát"],
            ["TVTK", "Tư vấn thiết kế"], ["NT", "Nhà thầu thi công xây dựng"], ["HSTK", "Hồ sơ thiết kế"],
            ["BPTC", "Biện pháp thi công"], ["ATLĐ, ATVSLĐ", "An toàn lao động; an toàn, vệ sinh lao động"],
            ["VSMT", "Vệ sinh môi trường"], ["PCCC và CNCH", "Phòng cháy, chữa cháy và cứu nạn, cứu hộ"],
            ["ITP", "Kế hoạch kiểm tra và thí nghiệm (Inspection and Test Plan)"], ["MIR / WIR", "Yêu cầu nghiệm thu vật liệu / công việc"],
            ["NCR", "Phiếu ghi nhận sự không phù hợp"], ["RFI", "Phiếu yêu cầu cung cấp, làm rõ thông tin"]]
    if D.level == "Cơ bản":
        rows = rows[:10]
    D.table([2200, 6826], ["Viết tắt", "Diễn giải"], rows, size=11)
    D.page_break()


def _phan_1(D, modules, LIB, PL):
    f, lv = D.info, D.level
    ADV = load("noi_dung_nang_cao.json")
    D.part("Thông tin chung và cơ sở thực hiện")
    if lv == "Nâng cao":
        D.h(2, "Mục tiêu của đề cương")
        D.snippet(ADV["muc_tieu"])
    D.h(2, "Thông tin dự án, gói thầu")
    pairs = [("Tên dự án", f.get("du_an")), ("Công trình", f.get("cong_trinh")), ("Gói thầu", f.get("goi_thau")),
             ("Địa điểm xây dựng", f.get("dia_diem")), ("Chủ đầu tư", f.get("chu_dau_tu")),
             ("Đại diện Chủ đầu tư / BQLDA", f.get("dai_dien_cdt")), ("Nguồn vốn", f.get("nguon_von")),
             ("Loại, cấp công trình", (f.get("loai_ct_ten") or f.get("loai_ct", "")) + (f"; cấp {f.get('cap_ct')}" if f.get("cap_ct") else "")),
             ("Hợp đồng TVGS", f.get("hop_dong")), ("Thời gian thực hiện", f"{f['thoi_gian']} tháng" if f.get("thoi_gian") else ""),
             ("Đơn vị TVGS", "Công ty Cổ phần TEXO Tư vấn và Đầu tư")]
    K.meta_table(D.d, [(a, b) for a, b in pairs if b], 9026, 0.32)
    D.p("", after=4)
    D.h(2, "Quy mô và giải pháp thiết kế chính")
    qm = []
    if f.get("so_tang"):
        qm.append(f"Số tầng nổi: {f['so_tang']} tầng" + (f"; số tầng hầm: {f['so_ham']} tầng." if int(f.get('so_ham') or 0) else "."))
    if f.get("dt_san"):
        qm.append(f"Tổng diện tích sàn xây dựng: {float(f['dt_san']):,.0f} m².".replace(",", "."))
    if f.get("chieu_cao"):
        qm.append(f"Chiều cao công trình: {f['chieu_cao']} m.")
    if f.get("mong"):
        qm.append("Giải pháp nền móng: " + ", ".join(f["mong"]).lower() + ".")
    if f.get("ket_cau"):
        qm.append("Giải pháp kết cấu: " + ", ".join(f["ket_cau"]) + ".")
    for line in (f.get("mo_ta_quy_mo") or "").splitlines():
        if line.strip():
            qm.append(line.strip().lstrip("-• "))
    D.bullets(qm or ["Theo hồ sơ thiết kế được phê duyệt."])
    ai = f.get("_ai") or {}
    if ai.get("dac_diem") or ai.get("trong_tam"):
        D.h(3, "Đặc điểm công trình và trọng tâm giám sát")
        D.bullets([x for x in ai.get("dac_diem", []) if isinstance(x, str)])
        rows = [[str(k + 1), t.get("cong_tac", ""), t.get("rui_ro", ""), t.get("bien_phap", "")]
                for k, t in enumerate(x for x in ai.get("trong_tam", []) if isinstance(x, dict))]
        if rows:
            D.table([600, 2200, 2900, 3326], ["TT", "Công tác", "Rủi ro chính", "Biện pháp giám sát, điểm kiểm soát"], rows, size=10.5, center_cols=(0,))
    if lv == "Nâng cao":
        D.h(3, "Các bên tham gia dự án")
        ben = f.get("cac_ben") or []
        rows = [["Chủ đầu tư", f.get("chu_dau_tu", "")], ["Quản lý dự án / Đại diện CĐT", f.get("dai_dien_cdt") or "Theo quyết định hoặc văn bản ủy quyền của Chủ đầu tư"],
                ["Tư vấn giám sát", "Công ty Cổ phần TEXO Tư vấn và Đầu tư"]] + [[a, b] for a, b in ben] + \
               [["Nhà thầu thi công", "Theo hợp đồng / quyết định lựa chọn của Chủ đầu tư"], ["Đơn vị thí nghiệm, quan trắc, kiểm định", "Theo hồ sơ năng lực và chấp thuận của Chủ đầu tư"]]
        D.table([700, 3000, 5326], ["TT", "Bên tham gia", "Tên / thông tin"], [[str(k + 1)] + r for k, r in enumerate(rows)], center_cols=(0,))
        D.h(3, "Hồ sơ và dữ liệu đầu vào phục vụ công tác giám sát")
        D.snippet(ADV["du_lieu_dau_vao"])

    D.h(2, "Phạm vi công việc của Tư vấn giám sát")
    pv = f.get("pham_vi") or ["Giám sát thi công xây dựng công trình", "Giám sát lắp đặt thiết bị vào công trình"]
    D.p("Tư vấn giám sát thay mặt Chủ đầu tư thực hiện giám sát thi công xây dựng theo hợp đồng TVGS đã ký kết, hồ sơ thiết kế được phê duyệt và quy định pháp luật hiện hành, gồm:")
    D.bullets(pv)
    D.bullets(["Giám sát chất lượng, khối lượng, tiến độ thi công xây dựng;",
               "Giám sát an toàn lao động, vệ sinh môi trường, phòng cháy chữa cháy trên công trường;",
               "Tham gia nghiệm thu công việc, giai đoạn, hạng mục và công trình hoàn thành theo quy định;",
               "Báo cáo định kỳ, đột xuất cho Chủ đầu tư và các công việc khác theo hợp đồng TVGS."])
    D.h(3, "Sản phẩm của gói thầu TVGS")
    D.bullets(["Đề cương tư vấn giám sát được Chủ đầu tư phê duyệt;",
               "Nhật ký giám sát, biên bản kiểm tra, biên bản nghiệm thu, xác nhận khối lượng theo phân cấp;",
               "Báo cáo tuần, báo cáo tháng, báo cáo giai đoạn và báo cáo đột xuất (nếu có);",
               "Báo cáo hoàn thành công tác giám sát thi công xây dựng theo mẫu quy định tại Nghị định số 207/2026/NĐ-CP."])
    if lv == "Nâng cao":
        D.snippet([x for x in ADV["pham_vi_tvgs"] if x[0] in ("p",) and x[1].startswith("(")])

    # ---- cơ sở thực hiện
    D.h(2, "Cơ sở thực hiện công tác giám sát")
    D.h(3, "Văn bản quy phạm pháp luật")
    vb = [x for x in PL["hien_hanh"] if lv == "Nâng cao" or x.get("mau") == "B"]
    rows = [[str(k + 1), x["ky_hieu"] + (f" ngày {x['ngay']}" if x.get("ngay") else ""), x["noi_dung"]] for k, x in enumerate(vb)]
    rows.append([str(len(rows) + 1), "Các văn bản quy phạm pháp luật khác", "Có liên quan, còn hiệu lực tại thời điểm thực hiện công việc"])
    D.table([700, 3900, 4426], ["TT", "Văn bản", "Nội dung áp dụng"], rows, size=10.5, center_cols=(0,))
    D.h(3, "Hồ sơ pháp lý của gói thầu")
    D.bullets([f"Hợp đồng tư vấn giám sát {f.get('hop_dong') or 'số … ngày …'} giữa {f.get('chu_dau_tu') or 'Chủ đầu tư'} và Công ty Cổ phần TEXO Tư vấn và Đầu tư, các phụ lục kèm theo;",
               "Hồ sơ thiết kế xây dựng triển khai sau thiết kế cơ sở được phê duyệt và được Chủ đầu tư kiểm tra, xác nhận trên bản vẽ; chỉ dẫn kỹ thuật, thuyết minh thiết kế;",
               "Hồ sơ mời thầu, hồ sơ dự thầu và hợp đồng thi công xây dựng giữa Chủ đầu tư và nhà thầu thi công; hợp đồng với nhà thầu phụ (nếu có);",
               "Quyết định thành lập đoàn Tư vấn giám sát của Công ty Cổ phần TEXO Tư vấn và Đầu tư;",
               "Những yêu cầu riêng của Chủ đầu tư quy định cho công trình."])
    D.h(3, "Quy chuẩn, tiêu chuẩn áp dụng")
    D.p("Áp dụng danh mục quy chuẩn, tiêu chuẩn được phê duyệt cùng hồ sơ thiết kế và chỉ dẫn kỹ thuật của dự án; quy chuẩn kỹ thuật quốc gia là bắt buộc trong phạm vi điều chỉnh. "
        "Tiêu chuẩn được dùng đúng phiên bản còn hiệu lực tại thời điểm áp dụng, trừ trường hợp hồ sơ thiết kế, hợp đồng quy định khác phù hợp pháp luật. "
        "Các tiêu chuẩn tham chiếu chính cho từng công tác như sau:")
    rows, seen = [], set()
    for m in modules:
        mm = LIB.get(m)
        if not mm:
            continue
        std = [s for s in mm["tieu_chuan"] if s not in seen][:8]
        seen.update(std)
        if std:
            rows.append([mm["ten"].replace("Giám sát ", "").replace("Kiểm tra ", "")[:1].upper() + mm["ten"].replace("Giám sát ", "").replace("Kiểm tra ", "")[1:], "; ".join(std)])
    rows.insert(0, ["An toàn trong thi công, PCCC", "QCVN 18:2021/BXD; QCVN 06:2022/BXD và sửa đổi (nếu áp dụng)"])
    D.table([700, 3300, 5026], ["TT", "Công tác", "Quy chuẩn, tiêu chuẩn tham chiếu"], [[str(k + 1)] + r for k, r in enumerate(rows)], size=10.5, center_cols=(0,))

    # ---- vai trò các bên
    D.h(2, "Vai trò, quyền và nghĩa vụ của các bên")
    D.h(3, "Chủ đầu tư")
    D.bullets(["Là bên quyết định cuối cùng về các vấn đề của dự án thuộc thẩm quyền của Chủ đầu tư;",
               "Quan hệ chính thức với các nhà thầu liên quan đến dự án bằng hợp đồng hoặc thỏa thuận trong khuôn khổ pháp luật;",
               "Cung cấp đầy đủ, kịp thời hồ sơ thiết kế, hợp đồng, tài liệu liên quan cho TVGS; xử lý kịp thời đề xuất của TVGS;",
               "Thông báo cho các bên liên quan về quyền và nghĩa vụ của kỹ sư TVGS;",
               "Yêu cầu TVGS thay đổi người giám sát khi không thực hiện đúng quy định; đình chỉ hoặc chấm dứt hợp đồng theo hợp đồng và pháp luật;",
               "Không thông đồng hoặc dùng ảnh hưởng để làm sai lệch kết quả giám sát; lưu trữ kết quả giám sát thi công xây dựng."])
    D.h(3, "Tư vấn giám sát TEXO")
    D.bullets(["Thực hiện đầy đủ nội dung công việc theo hợp đồng TVGS, đề cương được phê duyệt và quy định của Nghị định số 207/2026/NĐ-CP;",
               "Nghiệm thu, xác nhận khi công việc đã thi công đúng thiết kế, đúng quy chuẩn, tiêu chuẩn và bảo đảm chất lượng; từ chối nghiệm thu bằng văn bản khi không đạt yêu cầu;",
               "Phát hiện, đề xuất với Chủ đầu tư những bất hợp lý của thiết kế để kịp thời điều chỉnh;",
               "Yêu cầu nhà thầu thực hiện đúng hợp đồng thi công; tạm dừng thi công khi phát hiện vi phạm nghiêm trọng về chất lượng, an toàn và báo cáo ngay Chủ đầu tư;",
               "Bảo lưu ý kiến đối với công việc giám sát do mình đảm nhận; từ chối yêu cầu bất hợp lý của các bên;",
               "Chịu trách nhiệm trước Chủ đầu tư và pháp luật về chất lượng công tác giám sát của mình."])
    if lv == "Nâng cao":
        D.snippet(ADV["trach_nhiem_phoi_hop"])
    D.h(3, "Tư vấn thiết kế và nhà thầu thi công")
    D.bullets(["Tư vấn thiết kế: giám sát tác giả; giải thích, làm rõ hồ sơ thiết kế; xử lý vướng mắc, điều chỉnh thiết kế theo yêu cầu của Chủ đầu tư; tham gia nghiệm thu theo quy định;",
               "Nhà thầu thi công: tự tổ chức quản lý chất lượng, tự kiểm tra trước khi đề nghị nghiệm thu; thi công đúng thiết kế, biện pháp được duyệt; bảo đảm ATLĐ, VSMT, PCCC; lập hồ sơ hoàn công, hồ sơ quản lý chất lượng."])


def _phan_2(D, modules):
    lv = D.level
    ADV = load("noi_dung_nang_cao.json")
    D.page_break()
    D.part("Nội dung chính của công tác giám sát")
    # ---------------- chất lượng
    D.h(2, "Giám sát chất lượng thi công xây dựng")
    D.p("Trước khi nhà thầu thi công từng hạng mục chính, TVGS kiểm tra tại hiện trường các điều kiện thi công theo hồ sơ trúng thầu, hợp đồng thi công và hồ sơ được Chủ đầu tư phê duyệt; "
        "trong quá trình thi công, TVGS giám sát thường xuyên, liên tục, có hệ thống và tổ chức nghiệm thu theo quy định.")
    D.h(3, "Kiểm tra điều kiện khởi công, công tác chuẩn bị thi công")
    D.bullets(["Kiểm tra việc bàn giao mặt bằng, mốc định vị, tim trục, cao độ giữa Chủ đầu tư và nhà thầu, có sự chứng kiến của TVGS;",
               "Kiểm tra giấy phép xây dựng (với công trình phải có giấy phép) hoặc văn bản thay thế theo Luật Xây dựng số 135/2025/QH15;",
               "Kiểm tra hồ sơ thiết kế được phê duyệt, có dấu xác nhận của Chủ đầu tư trên bản vẽ; hợp đồng thi công xây dựng;",
               "Kiểm tra kế hoạch tổng hợp về an toàn, biện pháp bảo đảm an toàn, vệ sinh môi trường, an toàn công trình lân cận do nhà thầu lập."])
    if lv == "Nâng cao":
        D.snippet(ADV["chuan_bi_thi_cong"])
    D.h(3, "Kiểm tra sự phù hợp năng lực của nhà thầu thi công")
    D.sub("Nhân lực")
    D.bullets(["Đối chiếu nhân sự chủ chốt với hồ sơ trúng thầu/hợp đồng; mọi thay đổi phải được Chủ đầu tư chấp thuận bằng văn bản;",
               "Kiểm tra sơ đồ tổ chức, quyết định thành lập ban chỉ huy công trường, bằng cấp, chứng chỉ hành nghề, điều kiện năng lực theo Nghị định số 212/2026/NĐ-CP;",
               "Kiểm tra chứng chỉ thợ (hàn, vận hành máy), thẻ an toàn, hợp đồng lao động và trang bị bảo hộ lao động của công nhân."])
    D.sub("Máy, thiết bị thi công")
    D.bullets(["Danh sách, số lượng, chủng loại thiết bị phù hợp hồ sơ trúng thầu và biện pháp thi công được duyệt;",
               "Catalogue, tình trạng kỹ thuật; kết quả kiểm định an toàn, giấy phép sử dụng đối với máy, thiết bị có yêu cầu nghiêm ngặt về an toàn lao động."])
    D.h(3, "Kiểm tra hệ thống quản lý chất lượng của nhà thầu")
    D.bullets(["Sơ đồ tổ chức, quyền và nghĩa vụ của bộ phận, cá nhân chịu trách nhiệm quản lý chất lượng của nhà thầu;",
               "Kế hoạch và quy trình kiểm soát chất lượng vật liệu, thiết bị, công tác thi công; kế hoạch thí nghiệm, kiểm định, quan trắc;",
               "Quy trình lập và quản lý hồ sơ, nhật ký thi công, báo cáo nội bộ và báo cáo Chủ đầu tư."])
    D.h(3, "Kiểm tra phòng thí nghiệm và cơ sở sản xuất vật liệu")
    D.bullets(["Phòng thí nghiệm chuyên ngành xây dựng phải đủ điều kiện năng lực theo quy định, phù hợp hồ sơ trúng thầu và được Chủ đầu tư chấp thuận;",
               "Cơ sở sản xuất vật liệu, cấu kiện, bê tông thương phẩm cung cấp cho công trình phải có giấy phép, chứng nhận hợp chuẩn, hợp quy theo quy định."])
    D.h(3, "Kiểm tra, nghiệm thu vật liệu, sản phẩm, thiết bị lắp đặt vào công trình")
    D.bullets(["Vật liệu, thiết bị phải được Chủ đầu tư phê duyệt chủng loại, nguồn gốc trước khi nhập về công trường;",
               "Kiểm tra chứng chỉ xuất xưởng, CO/CQ, chứng nhận hợp quy theo Nghị định số 209/2026/NĐ-CP và Thông tư số 41/2026/TT-BXD; lấy mẫu thí nghiệm theo tần suất quy định;",
               "Khi nghi ngờ kết quả thí nghiệm, đề nghị Chủ đầu tư yêu cầu kiểm tra lại tại phòng thí nghiệm độc lập; không cho phép sử dụng vật liệu không đạt yêu cầu."])
    D.h(3, "Kiểm tra biện pháp thi công và bản vẽ shopdrawing")
    D.bullets(["Kiểm tra biện pháp thi công chi tiết so với hồ sơ trúng thầu; biện pháp đặc biệt hoặc biện pháp đòi hỏi thiết kế về an toàn chịu lực phải được thiết kế, thẩm tra riêng;",
               "Kiểm tra bản vẽ triển khai thi công, chế tạo, lắp đặt (shopdrawing) trên cơ sở thiết kế được duyệt, chỉ dẫn kỹ thuật và vật liệu được duyệt trước khi trình Chủ đầu tư; "
               "khi có sai khác, thiết kế được duyệt là căn cứ thực hiện."])
    D.h(3, "Giám sát thường xuyên, nhật ký thi công")
    D.bullets(["TVGS có mặt tại hiện trường theo tiến độ được duyệt; kiểm tra, giám sát liên tục, có hệ thống quá trình thi công; kết quả kiểm tra được ghi vào nhật ký hoặc biên bản kiểm tra;",
               "Nhật ký thi công được đánh số trang, đóng dấu giáp lai, cập nhật hằng ngày; là tài liệu gốc trao đổi thông tin giữa các bên."])
    D.h(3, "Kiểm tra, xác nhận bản vẽ hoàn công")
    D.bullets(["Nhà thầu lập bản vẽ hoàn công bộ phận, hạng mục, công trình trên cơ sở thiết kế được duyệt; bộ phận bị che khuất phải được nghiệm thu và vẽ hoàn công trước khi thi công tiếp;",
               "Kích thước, thông số thay đổi so với thiết kế được ghi bên cạnh trị số cũ và phải được Chủ đầu tư xác nhận; TVGS trưởng ký xác nhận bản vẽ hoàn công."])
    D.h(3, "Tổ chức nghiệm thu")
    D.snippet(ADV["nghiem_thu_2026"])
    D.h(3, "Hồ sơ nghiệm thu và hồ sơ hoàn thành công trình")
    D.bullets(["Tập hợp, kiểm tra tài liệu phục vụ nghiệm thu công việc, bộ phận, giai đoạn, hạng mục và công trình;",
               "Danh mục hồ sơ hoàn thành công trình lập theo Phụ lục của Nghị định số 207/2026/NĐ-CP: phần hồ sơ pháp lý do Chủ đầu tư thực hiện, TVGS phối hợp nhắc nhở; phần hồ sơ quản lý chất lượng do nhà thầu lập, TVGS kiểm tra, xác nhận;",
               "Hồ sơ được số hóa và lưu trữ trên hệ thống quản lý hồ sơ điện tử của Công ty Cổ phần TEXO Tư vấn và Đầu tư."])
    D.h(3, "Phát hiện sai sót, bất hợp lý của thiết kế")
    D.p("Khi phát hiện sai sót, bất hợp lý của hồ sơ thiết kế, TVGS báo cáo Chủ đầu tư để yêu cầu tư vấn thiết kế xem xét, điều chỉnh; TVGS không tự ý thay đổi thiết kế.")
    D.h(3, "Kiểm định lại chất lượng khi có nghi ngờ")
    D.p("Khi nghi ngờ chất lượng vật liệu, bộ phận, hạng mục công trình, TVGS đề nghị Chủ đầu tư yêu cầu nhà thầu kiểm định lại dưới sự chứng kiến của TVGS tại đơn vị thí nghiệm, kiểm định đủ điều kiện năng lực.")
    D.h(3, "Phối hợp giải quyết vướng mắc, phát sinh")
    D.p("Chủ đầu tư chủ trì giải quyết vướng mắc, phát sinh; các bên đưa ra ý kiến; Chủ đầu tư quyết định cuối cùng trong thẩm quyền. TVGS chủ động đề xuất giải pháp kỹ thuật để không làm gián đoạn thi công.")
    if lv == "Nâng cao":
        D.h(3, "Kế hoạch nghiệm thu và kiểm tra theo giai đoạn")
        D.snippet(ADV["ke_hoach_nghiem_thu"])
        D.h(3, "Các quy trình nghiệm thu tiêu biểu và xử lý không phù hợp (NCR)")
        D.snippet(ADV["quy_trinh_nghiem_thu"])
        D.h(3, "Quản lý thay đổi thiết kế, RFI và xung đột bộ môn")
        D.snippet(ADV["thay_doi_thiet_ke"])

    # ---------------- khối lượng
    D.h(2, "Giám sát khối lượng")
    if lv == "Nâng cao":
        D.h(3, "Nguyên tắc, quy trình và thời hạn xác nhận khối lượng")
        D.snippet(ADV["khoi_luong"])
    else:
        D.h(3, "Nguyên tắc xác nhận khối lượng theo hồ sơ thiết kế")
        D.bullets(["Khối lượng được xác nhận phải có đủ hồ sơ pháp lý, hồ sơ kỹ thuật; là khối lượng thực tế đã hoàn thành, đã được nghiệm thu đạt yêu cầu;",
                   "Khối lượng thực tế thi công do nhà thầu lập trên cơ sở bản vẽ hoàn công được các bên xác nhận; TVGS kiểm tra, xác nhận;",
                   "Khối lượng chênh lệch do thiết kế tính thừa, thiếu: Chủ đầu tư quyết định thanh toán căn cứ hợp đồng và xác nhận của tư vấn thiết kế."])
    D.h(3, "Khối lượng phát sinh so với hồ sơ thiết kế")
    D.bullets(["Khối lượng phát sinh tăng, giảm phải có thiết kế bổ sung hoặc chỉ thị được Chủ đầu tư phê duyệt đúng thẩm quyền; TVGS xác nhận trên cơ sở thực tế thi công;",
               "Chỉ thị công trường không tự xác lập đơn giá, giá trị thanh toán; việc thanh toán thực hiện theo hợp đồng và phụ lục hợp đồng."])
    D.h(3, "Khối lượng sửa đổi và khối lượng thi công khác")
    D.bullets(["Sửa đổi do thiết kế hoặc do Chủ đầu tư thay đổi chủng loại vật liệu được xác nhận theo nguyên tắc như khối lượng phát sinh;",
               "Khối lượng lán trại, văn phòng công trường, biện pháp thi công đặc biệt: TVGS chỉ xác nhận khi có văn bản phê duyệt thiết kế, dự toán của Chủ đầu tư;",
               "Khối lượng thuộc chi phí trực tiếp khác, phục vụ thi công: TVGS không xác nhận."])

    # ---------------- tiến độ
    D.h(2, "Giám sát tiến độ thi công")
    if lv == "Nâng cao":
        D.h(3, "Kiểm soát tiến độ, cảnh báo sớm và biện pháp khắc phục chậm tiến độ")
        D.snippet(ADV["tien_do"])
    else:
        D.h(3, "Theo dõi, kiểm tra tiến độ thi công")
        D.bullets(["Theo dõi tiến độ căn cứ tiến độ thi công chi tiết do nhà thầu lập, được Chủ đầu tư phê duyệt; đối chiếu thực tế hằng tuần;",
                   "Thường xuyên cảnh báo nhà thầu và báo cáo Chủ đầu tư khi có nguy cơ chậm tiến độ, nêu rõ nguyên nhân;",
                   ])
        D.h(3, "Biện pháp khắc phục chậm tiến độ, đẩy nhanh tiến độ")
        D.bullets(["Yêu cầu nhà thầu lập biện pháp khắc phục, đẩy nhanh tiến độ (bổ sung nhân lực, thiết bị, tăng ca, điều chỉnh trình tự); TVGS kiểm tra tính khả thi và theo dõi thực hiện;",
                   "Trường hợp tổng tiến độ bị kéo dài, kiến nghị Chủ đầu tư báo cáo người quyết định đầu tư xem xét điều chỉnh."])

    # ---------------- an toàn
    D.h(2, "Giám sát an toàn lao động, vệ sinh môi trường, phòng cháy chữa cháy")
    D.bullets(["Nhà thầu thi công chịu trách nhiệm chính về ATLĐ, VSMT, PCCC trên công trường; TVGS kiểm tra, giám sát việc thực hiện và báo cáo Chủ đầu tư;",
               "Kiểm tra kế hoạch tổng hợp về an toàn, biện pháp an toàn chung và biện pháp an toàn cho công việc có nguy cơ cao theo QCVN 18:2021/BXD;",
               "Kiểm tra hồ sơ huấn luyện ATVSLĐ, thẻ an toàn, hợp đồng lao động; kiểm định máy, thiết bị có yêu cầu nghiêm ngặt về an toàn;",
               "Kiểm tra trang bị bảo hộ lao động, hệ thống biển báo, rào chắn, lưới chống rơi, điện thi công, phương tiện PCCC;",
               "Kiểm tra biện pháp bảo vệ môi trường: bụi, tiếng ồn, nước thải, chất thải rắn; biện pháp bảo vệ công trình lân cận;",
               "Khi phát hiện nguy cơ mất an toàn nghiêm trọng, TVGS yêu cầu dừng ngay phần việc liên quan và báo cáo Chủ đầu tư."])
    if lv == "Nâng cao":
        D.snippet(ADV["atld_rui_ro"])
        D.h(3, "Quy trình đình chỉ khẩn cấp và cho phép thi công lại")
        D.snippet(ADV["dinh_chi"])

    # ---------------- báo cáo
    D.h(2, "Chế độ báo cáo và tổ chức các cuộc họp")
    D.h(3, "Chế độ báo cáo")
    D.bullets(["Báo cáo tuần, báo cáo tháng về tình hình chất lượng, khối lượng, tiến độ, an toàn và kiến nghị;",
               "Báo cáo giai đoạn, báo cáo hoàn thành hạng mục; báo cáo đột xuất khi có sự cố hoặc vấn đề phát sinh;",
               "Báo cáo hoàn thành công tác giám sát thi công xây dựng theo mẫu quy định tại Nghị định số 207/2026/NĐ-CP."])
    D.h(3, "Thời hạn nộp báo cáo")
    if lv == "Nâng cao":
        D.snippet(ADV["bao_cao"])
    else:
        D.table([700, 3600, 4726], ["TT", "Loại báo cáo", "Thời hạn nộp"], [
            ["1", "Báo cáo đầu kỳ (kế hoạch huy động nhân sự, kiểm tra điều kiện khởi công)", "Sau 05 ngày kể từ khi tiếp nhận công việc"],
            ["2", "Báo cáo đột xuất, sự cố", "Chậm nhất 24 giờ kể từ khi phát sinh"],
            ["3", "Báo cáo tuần", "Chiều thứ Hai tuần kế tiếp hoặc trước buổi giao ban định kỳ"],
            ["4", "Báo cáo tháng", "Ngày 02 của tháng kế tiếp"],
            ["5", "Báo cáo giai đoạn / hoàn thành hạng mục", "Sau khi nghiệm thu giai đoạn, hạng mục"],
            ["6", "Báo cáo hoàn thành công tác giám sát", "Tối đa 15 ngày sau khi hoàn thành công trình"]], center_cols=(0,))
    D.h(3, "Tổ chức các cuộc họp")
    D.bullets(["Các cuộc họp do Chủ đầu tư tổ chức; TVGS và các nhà thầu tham dự, cho ý kiến;",
               "Họp giao ban định kỳ hằng tuần/tháng tại công trường do TVGS trưởng tham dự; lãnh đạo Công ty tham dự khi Chủ đầu tư yêu cầu;",
               "Họp đột xuất khi có nguy cơ ảnh hưởng chất lượng, tiến độ, an toàn; giấy mời gửi trước ít nhất 03 ngày (trừ trường hợp khẩn cấp)."])
    if lv == "Nâng cao":
        D.h(3, "Trao đổi thông tin và quản lý hồ sơ, tài liệu")
        D.snippet(ADV["giao_tiep"])
        D.snippet(ADV["quan_ly_ho_so"])
        D.snippet(ADV["tai_lieu_bat_buoc"])

    # ---------------- tổ chức
    D.h(2, "Tổ chức thực hiện của đoàn Tư vấn giám sát")
    D.h(3, "Nguyên tắc chung")
    D.bullets(["Đề cương sau khi được Chủ đầu tư phê duyệt là tài liệu điều chỉnh việc thực hiện hợp đồng TVGS giữa các bên liên quan;",
               "Công ty Cổ phần TEXO Tư vấn và Đầu tư thành lập đoàn TVGS bằng quyết định sau khi hợp đồng được ký kết; tiến độ huy động nhân sự theo tiến độ thi công chi tiết."])
    D.h(3, "Mô hình đoàn Tư vấn giám sát và thành phần nhân sự")
    staff = D.info.get("_staff") or []
    D.table([600, 2400, 1900, 4126], ["TT", "Chức danh", "Họ và tên", "Chuyên môn, nhiệm vụ chính"],
            [[str(k + 1), s["chuc_danh"], s.get("ho_ten") or "", f"{s.get('chuyen_mon', '')}. {s.get('nhiem_vu', '')}".strip(". ")] for k, s in enumerate(staff)],
            size=10.5, center_cols=(0,))
    D.h(3, "Nhiệm vụ của TVGS trưởng và các kỹ sư giám sát")
    D.bullets(["TVGS trưởng: thay mặt Tổng Giám đốc quản lý toàn diện công tác giám sát của đoàn; điều phối, phân công kỹ sư; trực tiếp làm việc với Chủ đầu tư;",
               "Kỹ sư TVGS: chịu sự phân công của TVGS trưởng; có mặt tại hiện trường theo giai đoạn thi công; báo cáo ngay TVGS trưởng khi phát hiện chất lượng không đạt."])
    if D.level == "Nâng cao":
        D.h(3, "Kế hoạch điều động, đánh giá và đào tạo nhân sự")
        D.snippet(ADV["dieu_dong_dao_tao"])
    D.h(3, "Quan hệ của đoàn TVGS với Công ty và các bên")
    D.bullets(["Công ty bố trí bộ phận văn phòng theo dõi thường xuyên diễn biến công trường, hỗ trợ kỹ thuật cho đoàn;",
               "Bộ phận quản lý kỹ thuật của Công ty hỗ trợ, kiểm soát việc thực hiện hợp đồng giám sát theo hệ thống quản lý chất lượng của Công ty;",
               "Với Chủ đầu tư: TVGS trưởng là đầu mối; kỹ sư TVGS là thành viên nghiệm thu theo phân công;",
               "Với nhà thầu thi công: giám sát chặt chẽ, đôn đốc báo cáo; có quyền yêu cầu tạm dừng khi phát hiện sai phạm nghiêm trọng (sau khi báo cáo Chủ đầu tư);",
               "Với tư vấn thiết kế: phối hợp thông qua Chủ đầu tư để hiểu đúng thiết kế, đề nghị làm rõ, ghi nhận thay đổi thiết kế vào nhật ký hoặc văn bản."])
    D.h(3, "Phân cấp trách nhiệm ký hồ sơ sản phẩm TVGS")
    D.table([700, 5200, 3126], ["TT", "Hồ sơ, tài liệu", "Người ký phía TVGS"], [
        ["1", "Biên bản bàn giao hồ sơ, tài liệu", "Kỹ sư TVGS"],
        ["2", "Biên bản bàn giao mặt bằng, mốc, công tác chuẩn bị khởi công", "TVGS trưởng"],
        ["3", "Biện pháp thi công, shopdrawing của nhà thầu (xác nhận đã kiểm tra)", "TVGS trưởng"],
        ["4", "Biên bản nghiệm thu vật liệu, công việc xây dựng", "Kỹ sư TVGS được phân công"],
        ["5", "Biên bản nghiệm thu bộ phận, giai đoạn thi công", "TVGS trưởng"],
        ["6", "Bản vẽ hoàn công", "TVGS trưởng"],
        ["7", "Xác nhận khối lượng thanh toán các giai đoạn", "TVGS trưởng"],
        ["8", "Biên bản nghiệm thu hoàn thành hạng mục, công trình đưa vào sử dụng", "Lãnh đạo Công ty/ người được ủy quyền (TVGS trưởng tham gia)"],
        ["9", "Xác nhận khối lượng quyết toán", "Lãnh đạo Công ty/ người được ủy quyền (TVGS trưởng tham gia)"],
        ["10", "Chỉ thị công trường, thư báo kỹ thuật, văn bản công trường", "TVGS trưởng / Kỹ sư TVGS"],
        ["11", "Báo cáo tuần, tháng, giai đoạn; báo cáo sự cố", "TVGS trưởng"],
        ["12", "Báo cáo hoàn thành công tác giám sát", "TVGS trưởng lập, lãnh đạo Công ty ký"]], size=10.5, center_cols=(0,))
    if D.level == "Nâng cao":
        D.h(3, "Ma trận trách nhiệm phối hợp các bên")
        D.snippet(ADV["ma_tran"])
        D.h(3, "Hệ thống đảm bảo chất lượng công tác giám sát và phòng ngừa sai sót")
        D.bullets(["Công ty áp dụng hệ thống quản lý chất lượng ISO 9001 cho công tác tư vấn giám sát; sản phẩm của đoàn được kiểm soát theo quy trình kiểm soát sản phẩm của Công ty;",
                   "Bộ phận quản lý kỹ thuật của Công ty kiểm tra định kỳ hồ sơ, báo cáo, biên bản của đoàn; kiểm tra đột xuất khi có dấu hiệu rủi ro;",
                   "Phòng ngừa sai sót chủ quan: phổ biến hồ sơ thiết kế, ITP, checklist cho từng kỹ sư trước khi thi công hạng mục mới; kiểm tra chéo trước khi ký;",
                   "Phòng ngừa sai sót khách quan: theo dõi thay đổi thiết kế, thời tiết, nguồn cung vật liệu; cảnh báo sớm và đề xuất giải pháp với Chủ đầu tư."])
        D.h(3, "Kế hoạch quản lý an toàn của đoàn TVGS")
        D.snippet(ADV["ke_hoach_an_toan"][:60])


def _module(D, m, idx):
    D.h(3, m["ten"])
    tcs = m.get("tieu_chuan_tham_chieu") or []
    if tcs:
        D.sub("Tiêu chuẩn, tài liệu tham chiếu")
        D.bullets(tcs[:10])
    rows = m.get("hang") or []
    if rows:
        widths = [3000, 6026]
        t = K.new_table(D.d, widths, rows=1)
        hdr = t.rows[0]
        K.no_split(hdr, header=True)
        for j, h in enumerate(["Nội dung / tiêu chuẩn / hồ sơ xác nhận", "Phương pháp giám sát, kiểm soát chất lượng"]):
            c = hdr.cells[j]
            K._tc_borders(c, K.GRID)
            K.shade(c, K.C["blue"])
            K.margins(c)
            K.cp(c, h, 10.5, b=True, c=K.C["white"], align=WD_ALIGN_PARAGRAPH.CENTER)
        group = None
        for r in rows:
            if r.get("nhom") and r["nhom"] != group:
                group = r["nhom"]
                row = t.add_row()
                K.no_split(row)
                a, b = row.cells
                cell = a.merge(b)
                K._tc_borders(cell, K.GRID)
                K.shade(cell, K.C["label"])
                K.margins(cell)
                K.cp(cell, group, 10.5, b=True, c=K.C["navy"])
            row = t.add_row()
            a, b = row.cells
            for c in (a, b):
                K._tc_borders(c, K.GRID)
                K.margins(c)
            K.shade(a, K.C["first"])
            K.cp(a, r.get("muc") or "", 10.5, b=True)
            if r.get("tc"):
                K.cp(a, "Tiêu chuẩn: " + "; ".join(r["tc"][:6]), 9.5, i=True, c=K.C["grey"], first=False)
            if r.get("hs"):
                K.cp(a, "Hồ sơ: " + "; ".join(r["hs"][:5]), 9.5, i=True, c=K.C["teal"], first=False)
            if r.get("diem"):
                K.cp(a, "Điểm kiểm soát: " + r["diem"], 9.5, b=True, c=K.C["orange"], first=False)
            first = True
            for line in r.get("pp") or []:
                txt = line if re.match(r"^([a-zđ]\)|\d+[\.\)])", line) else "- " + line
                K.cp(b, txt, 10.5, first=first)
                first = False
        D.p("", after=4)
    if D.level == "Nâng cao":
        for ck in (m.get("checklist") or [])[:4]:
            rows = ck["rows"]
            if len(rows) < 3:
                continue
            ncol = max(len(r) for r in rows)
            head = [" ".join(c) for c in rows[0]] + [""] * (ncol - len(rows[0]))
            # bỏ hàng tiêu đề thứ 2 trùng (NTTC/TVGS/BQL)
            body = rows[1:]
            if body and " ".join(" ".join(c) for c in body[0]).count("TVGS") and len(body[0]) == ncol:
                head = [" ".join(c) for c in body[0]]
                body = body[1:]
            body = [[" ".join(c) for c in r] + [""] * (ncol - len(r)) for r in body][:45]
            D.h(4, "Phiếu kiểm tra: " + (ck.get("ten") or "").rstrip(":"))
            if ncol <= 3:
                widths = [3200, 3826, 2000][:ncol]
            else:
                widths = [2600, 3626] + [int(2800 / (ncol - 2))] * (ncol - 2)
            widths[-1] += 9026 - sum(widths)
            D.table(widths, head, body, size=9.5)
        for bp in (m.get("bang_phu") or [])[:3]:
            rows = [[" ".join(c) for c in r] for r in bp["rows"]]
            ncol = max(len(r) for r in rows)
            rows = [r + [""] * (ncol - len(r)) for r in rows]
            if len(rows) < 2:
                continue
            w = [9026 // ncol] * ncol
            w[-1] += 9026 - sum(w)
            D.sub(bp.get("ten") or "Bảng tham chiếu")
            D.table(w, rows[0], rows[1:], size=9.5)


def _phan_3(D, modules, LIB):
    D.page_break()
    D.part("Đề cương giám sát chi tiết các công tác")
    D.h(2, "Nguyên tắc áp dụng")
    D.bullets(["Đề cương chi tiết nêu nội dung, phương pháp giám sát và hồ sơ xác nhận cho từng công tác chính của công trình; thông số kỹ thuật cụ thể áp dụng theo hồ sơ thiết kế, chỉ dẫn kỹ thuật và tiêu chuẩn được phê duyệt;",
               "Công tác chưa nêu trong đề cương được giám sát theo hồ sơ thiết kế, biện pháp thi công được duyệt và tiêu chuẩn thi công, nghiệm thu tương ứng;",
               "Mỗi công việc chỉ được chuyển bước khi đã được nghiệm thu đạt yêu cầu."])
    if any(r.get("diem") for mid in modules for r in (LIB.get(mid) or {}).get("hang", [])):
        D.bullets(["Điểm kiểm soát ghi trong bảng: HP (Hold Point) – điểm dừng bắt buộc, chỉ chuyển bước khi TVGS chấp thuận; "
                   "WP (Witness Point) – điểm chứng kiến, nhà thầu báo trước để TVGS chứng kiến; RP (Review Point) – điểm xem xét hồ sơ."])
    groups = []
    for mid in modules:
        m = LIB.get(mid)
        if m and m["nhom"] not in groups:
            groups.append(m["nhom"])
    ten_nhom = {"Chung": "Kiểm tra đầu vào và công tác chung", "Nền móng": "Giám sát phần nền móng, tầng hầm", "Kết cấu": "Giám sát phần kết cấu",
                "Hoàn thiện": "Giám sát phần hoàn thiện", "Cơ điện (MEP)": "Giám sát lắp đặt hệ thống cơ điện (MEP)", "Hạ tầng": "Giám sát hạ tầng kỹ thuật ngoài nhà",
                "Chuyên đề": "Giám sát công tác chuyên đề", "Công trình xử lý nước thải": "Giám sát bể, đường ống, thiết bị công nghệ và chạy thử"}
    order = ["Chung", "Nền móng", "Kết cấu", "Công trình xử lý nước thải", "Hoàn thiện", "Cơ điện (MEP)", "Hạ tầng", "Chuyên đề"]
    for g in [x for x in order if x in groups]:
        D.h(2, ten_nhom[g])
        for k, mid in enumerate(modules):
            m = LIB.get(mid)
            if m and m["nhom"] == g:
                _module(D, m, k)


def _phan_4(D):
    D.page_break()
    D.part("Biểu mẫu hồ sơ")
    D.p("Áp dụng biểu mẫu theo Nghị định số 207/2026/NĐ-CP, Thông tư số 32/2026/TT-BXD và biểu mẫu dự án được Chủ đầu tư ban hành. Danh mục biểu mẫu chính:")
    rows = [["Phiếu yêu cầu nghiệm thu (vật liệu, công việc, giai đoạn)", "Nhà thầu lập, TVGS tiếp nhận"],
            ["Biên bản nghiệm thu vật liệu, sản phẩm, thiết bị đầu vào", "Kỹ sư TVGS"],
            ["Biên bản lấy mẫu thí nghiệm", "Kỹ sư TVGS chứng kiến"],
            ["Biên bản nghiệm thu công việc xây dựng", "Kỹ sư TVGS"],
            ["Biên bản kiểm tra hồ sơ nghiệm thu giai đoạn / hoàn thành", "TVGS trưởng"],
            ["Biên bản nghiệm thu bộ phận, giai đoạn thi công xây dựng", "TVGS trưởng"],
            ["Biên bản nghiệm thu hoàn thành hạng mục, công trình", "Lãnh đạo Công ty / người được ủy quyền"],
            ["Nhật ký thi công xây dựng công trình", "Nhà thầu lập, TVGS ghi ý kiến"],
            ["Chỉ thị công trường / thư báo kỹ thuật", "TVGS"],
            ["Biên bản kiểm tra ATLĐ, VSMT, PCCC", "Giám sát ATLĐ"],
            ["Báo cáo tuần, tháng của TVGS", "TVGS trưởng"],
            ["Báo cáo hoàn thành công tác giám sát", "TVGS trưởng lập"]]
    if D.level == "Nâng cao":
        rows[8:8] = [["Phiếu ghi nhận sự không phù hợp (NCR)", "TVGS"], ["Phiếu yêu cầu cung cấp, làm rõ thông tin (RFI)", "Nhà thầu / TVGS"],
                     ["Danh mục theo dõi hồ sơ trình duyệt (shopdrawing, vật liệu, biện pháp)", "Kỹ sư QA-QC"],
                     ["Phiếu kiểm tra (checklist) theo từng công tác – Phần III", "Kỹ sư TVGS"]]
    D.table([700, 5300, 3026], ["TT", "Biểu mẫu", "Người lập / ký"], [[str(k + 1)] + r for k, r in enumerate(rows)], size=10.5, center_cols=(0,))


def generate(info, level, modules, staff, out_path):
    """info: dict thông tin dự án; level: 'Cơ bản'|'Nâng cao'; modules: list id; staff: list dict."""
    LIB = {m["id"]: m for m in load("thu_vien_hang_muc.json")["hang_muc"]}
    PL = load("phap_ly.json")
    info = dict(info)
    info["_staff"] = staff
    D = Doc(info, level)
    d = D.d
    s0 = d.sections[0]
    s0.page_width, s0.page_height = Twips(11906), Twips(16838)
    s0.left_margin = s0.right_margin = Twips(1440)
    s0.top_margin, s0.bottom_margin = Twips(1300), Twips(1080)
    _cover(D, approval=False)
    _cover(D, approval=True)
    _toc(D)
    D.page_break()
    _viet_tat(D)
    # phần thân: section mới có header/footer TEXO
    K.ORG["dept"] = "TƯ VẤN GIÁM SÁT THI CÔNG XÂY DỰNG"   # không nêu đơn vị nội bộ trên tài liệu phát hành
    ten_ngan = info.get("ten_ngan") or (info.get("cong_trinh") or info.get("du_an") or "")[:34]
    sec = d.add_section(WD_SECTION.CONTINUOUS)
    K.setup_section(sec, "A4P", "ĐỀ CƯƠNG TƯ VẤN GIÁM SÁT", ten_ngan[:40],
                    f"Đề cương TVGS · {(info.get('goi_thau') or '')[:60]} · Lần ban hành 01")
    # bỏ header/footer ở section bìa
    s0.header.is_linked_to_previous = False
    s0.footer.is_linked_to_previous = False
    # di chuyển ngắt section lên trước nội dung: python-docx thêm sectPr vào đoạn cuối hiện tại → ổn
    _phan_1(D, modules, LIB, PL)
    _phan_2(D, modules)
    _phan_3(D, modules, LIB)
    _phan_4(D)
    _fill_toc(D)
    d.core_properties.title = "Đề cương TVGS – " + (info.get("du_an") or "")
    d.core_properties.author = "Công ty Cổ phần TEXO Tư vấn và Đầu tư"
    d.save(out_path)
    return out_path
